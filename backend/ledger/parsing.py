"""Streaming readers and strict normalization for financial source files."""

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

# Aliases are ordered: the first matching name wins deterministically.
ALIASES = {
    "reference": ("reference", "transaction_reference", "invoice_id"),
    "amount": ("amount", "transaction_amount"),
    "currency": ("currency", "currency_code"),
    "timestamp": ("timestamp", "transaction_date", "date"),
    "counterparty": ("counterparty", "customer", "merchant"),
    "direction": ("direction", "entry_type"),
    "description": ("description", "memo", "narration"),
    "source_id": ("source_id", "event_id", "transaction_id"),
}
REQUIRED_HEADERS = ("reference", "amount", "currency", "timestamp")
MAX_ROWS = 250_000


class FileValidationError(ValueError):
    """The entire file cannot be read safely or has an invalid schema."""


class RowValidationError(ValueError):
    """One row is invalid; other valid rows can still be imported."""


def normalize_reference(value: str) -> str:
    # Normalize case and whitespace without deleting meaningful punctuation.
    return " ".join(value.strip().upper().split())


def header_key(value) -> str:
    return "_".join(str(value or "").strip().lower().split())


def validate_headers(headers) -> list[str]:
    keys = [header_key(value) for value in headers]
    if not keys or any(not key for key in keys) or len(set(keys)) != len(keys):
        raise FileValidationError("Headers must be nonempty and unique.")
    missing = [
        name for name in REQUIRED_HEADERS
        if not any(alias in keys for alias in ALIASES[name])
    ]
    if missing:
        raise FileValidationError(f"Missing headers: {', '.join(missing)}.")
    return keys


def check_excel_archive(file_object) -> None:
    # Reject unexpectedly large expanded XLSX content before opening the XML.
    try:
        with zipfile.ZipFile(file_object) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 200_000_000:
                raise FileValidationError("The expanded Excel workbook is too large.")
    except zipfile.BadZipFile as exc:
        raise FileValidationError("This is not a valid XLSX file.") from exc
    finally:
        file_object.seek(0)


def raw_rows(file_object, filename: str):
    """Yield (source row number, dictionary) without materializing the file."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        stream = io.TextIOWrapper(file_object, encoding="utf-8-sig", newline="")
        try:
            reader = csv.reader(stream)
            headers = validate_headers(next(reader, []))
            for number, values in enumerate(reader, start=2):
                if number > MAX_ROWS + 1:
                    raise FileValidationError(f"A file cannot exceed {MAX_ROWS:,} rows.")
                if not any(str(value).strip() for value in values):
                    continue
                # A malformed row is represented explicitly, not silently truncated.
                if len(values) != len(headers):
                    yield number, {"__error__": "Column count differs from header."}
                else:
                    yield number, dict(zip(headers, values, strict=True))
        except UnicodeDecodeError as exc:
            raise FileValidationError("CSV files must be UTF-8 encoded.") from exc
        finally:
            # Detach so closing TextIOWrapper does not close Django's file handle.
            stream.detach()
    elif suffix == ".xlsx":
        check_excel_archive(file_object)
        workbook = load_workbook(file_object, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = validate_headers(next(rows, []))
            for number, values in enumerate(rows, start=2):
                if number > MAX_ROWS + 1:
                    raise FileValidationError(f"A file cannot exceed {MAX_ROWS:,} rows.")
                if any(value is not None for value in values):
                    yield number, dict(zip(headers, values, strict=True))
        finally:
            workbook.close()
    elif suffix == ".xls":
        # xlrd handles legacy binary spreadsheets. Unlike XLSX read-only mode,
        # its workbook is materialized, so the upload serializer caps XLS at 10 MB.
        import xlrd
        workbook = xlrd.open_workbook(file_contents=file_object.read(), on_demand=True)
        try:
            sheet = workbook.sheet_by_index(0)
            if not sheet.nrows:
                raise FileValidationError("The workbook is empty.")
            if sheet.nrows > MAX_ROWS + 1:
                raise FileValidationError(f"A file cannot exceed {MAX_ROWS:,} rows.")
            headers = validate_headers(sheet.row_values(0))
            for index in range(1, sheet.nrows):
                values = []
                for cell in sheet.row(index):
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    values.append(value)
                if any(value != "" for value in values):
                    yield index + 1, dict(zip(headers, values, strict=True))
        finally:
            workbook.release_resources()
    else:
        raise FileValidationError("Choose a CSV, XLSX, or XLS file.")


def iter_batch_rows(batch):
    with batch.stored_file.open("rb") as file_object:
        yield from raw_rows(file_object, batch.original_filename)


def count_batch_rows(batch) -> int:
    """Count input rows before importing without retaining their contents."""
    return sum(1 for _ in iter_batch_rows(batch))


def calculate_file_sha256(uploaded_file) -> str:
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)  # Rewind before Django saves the uploaded content.
    return digest.hexdigest()


@dataclass(frozen=True)
class NormalizedTransaction:
    reference: str
    amount: Decimal
    currency: str
    timestamp: datetime
    counterparty: str
    direction: str
    description: str
    metadata: dict
    source_id: str = ""


def normalize_row(raw: dict) -> NormalizedTransaction:
    if "__error__" in raw:
        raise RowValidationError(raw["__error__"])
    row = {}
    for name, aliases in ALIASES.items():
        row[name] = next((raw[key] for key in aliases if key in raw), "")
    if any(row[name] is None or str(row[name]).strip() == "" for name in REQUIRED_HEADERS):
        raise RowValidationError("Reference, amount, currency and timestamp are required.")
    reference = str(row["reference"]).strip()
    if len(reference) > 255:
        raise RowValidationError("Reference cannot exceed 255 characters.")
    try:
        amount = Decimal(str(row["amount"]))
    except (InvalidOperation, ValueError) as exc:
        raise RowValidationError("Amount is not a valid decimal.") from exc
    if not amount.is_finite() or not Decimal("0") <= amount < Decimal("1000000000000000"):
        raise RowValidationError("Amount must be finite, nonnegative and fit DECIMAL(19,4).")
    # Reject lost precision rather than silently rounding financial data.
    if amount != amount.quantize(Decimal("0.0001")):
        raise RowValidationError("Amount cannot exceed four decimal places.")
    amount = amount.quantize(Decimal("0.0001"))
    currency = str(row["currency"]).strip().upper()
    if len(currency) != 3 or not currency.isascii() or not currency.isalpha():
        raise RowValidationError("Currency must be a three-letter code, such as USD.")
    value = row["timestamp"]
    try:
        if isinstance(value, datetime):
            timestamp = value
        elif isinstance(value, date):
            timestamp = datetime.combine(value, datetime.min.time())
        else:
            timestamp = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=timezone.utc)
        timestamp = timestamp.astimezone(timezone.utc)
    except (ValueError, OverflowError) as exc:
        raise RowValidationError("Timestamp must use ISO 8601 format.") from exc
    direction = str(row["direction"] or "CREDIT").strip().upper()
    if direction not in {"CREDIT", "DEBIT"}:
        raise RowValidationError("Direction must be CREDIT or DEBIT.")
    counterparty = str(row["counterparty"] or "").strip()
    if len(counterparty) > 255:
        raise RowValidationError("Counterparty cannot exceed 255 characters.")
    # JSON-safe metadata preserves extra columns, including Excel date values.
    known = {alias for aliases in ALIASES.values() for alias in aliases}
    metadata = {key: str(value) for key, value in raw.items() if key not in known}
    return NormalizedTransaction(
        reference, amount, currency, timestamp, counterparty, direction,
        str(row["description"] or ""), metadata, str(row["source_id"] or ""),
    )


def build_transaction_hash(row: NormalizedTransaction) -> str:
    # Currency, direction and a native event ID avoid false identity collisions.
    # Tenant, source type and account are scoped by the database constraint.
    values = [
        normalize_reference(row.reference), f"{row.amount:.4f}",
        row.timestamp.isoformat(), row.currency, row.direction, row.source_id,
    ]
    canonical = json.dumps(values, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()
